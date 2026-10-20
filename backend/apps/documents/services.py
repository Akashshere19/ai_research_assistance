
import csv
from pathlib import Path
import pymupdf
# import fitz
import pandas as pd
import pytesseract
from PIL import Image
from docx import Document as WordDocument


def extract_document_content(file_path, file_type):
    file_type = file_type.lower().lstrip(".")

    if file_type == "pdf":
        return extract_pdf(file_path)

    if file_type in {"png", "jpg", "jpeg"}:
        return pytesseract.image_to_string(Image.open(file_path))

    if file_type in {"txt", "md"}:
        return Path(file_path).read_text(
            encoding="utf-8",
            errors="replace",
        )

    if file_type == "csv":
        return extract_csv(file_path)

    if file_type == "xlsx":
        return extract_excel(file_path)
    
    if file_type == "docx":
        return extract_docx(file_path)

    raise ValueError(f"Unsupported file type: {file_type}")


def extract_pdf(file_path):
    pages = []

    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text().strip()

            # OCR fallback for scanned pages.
            if not text:
                pixmap = page.get_pixmap(dpi=200)
                image = Image.frombytes(
                    "RGB",
                    [pixmap.width, pixmap.height],
                    pixmap.samples,
                )
                text = pytesseract.image_to_string(image).strip()

            pages.append(
                f"[Page {page_number}]\n{text}"
            )

    return "\n\n".join(pages)


def extract_csv(file_path):
    rows = []

    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows.extend(
            " | ".join(cell.strip() for cell in row)
            for row in reader
        )

    return "\n".join(rows)


def extract_excel(file_path):
    sheets = pd.read_excel(
        file_path,
        sheet_name=None,
        header=None,
        dtype=str,
    )

    output = []

    for sheet_name, dataframe in sheets.items():
        output.append(f"[Sheet: {sheet_name}]")

        dataframe = dataframe.fillna("")
        for row in dataframe.values.tolist():
            output.append(
                " | ".join(str(cell).strip() for cell in row)
            )

    return "\n".join(output)



def extract_docx(file_path):
    document = WordDocument(file_path)
    content = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            content.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        content.append(f"\n[Table {table_index}]")

        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            content.append(" | ".join(cells))

    return "\n".join(content)
